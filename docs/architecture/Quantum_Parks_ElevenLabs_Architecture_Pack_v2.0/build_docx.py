from pathlib import Path
import re, csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image

ROOT=Path(__file__).parent
MD=ROOT/'ARCHITECTURE.md'
OUT=ROOT/'Quantum_Parks_ElevenLabs_Architecture_Pack_v2.0.docx'

NAVY='163A5F'; BLUE='2E6F9E'; TEAL='1C7A72'; LIGHT='EAF2F8'; LIGHT2='EEF7F4'; GRAY='5C6B73'; BORDER='CBD7E1'; WHITE='FFFFFF'; DARK='18232D'; AMBER='F6E9CF'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn(f'w:{m}'))
        if node is None:
            node=OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def set_run_font(run, name='Liberation Sans', size=None, bold=None, color=None, italic=None):
    run.font.name=name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),name)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor.from_string(color)

def add_bottom_border(paragraph, color=BORDER, size='8'):
    pPr=paragraph._p.get_or_add_pPr(); pBdr=pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr=OxmlElement('w:pBdr'); pPr.append(pBdr)
    bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),size); bottom.set(qn('w:space'),'6'); bottom.set(qn('w:color'),color); pBdr.append(bottom)

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run('Page '); set_run_font(run,size=8,color=GRAY)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
    paragraph._p.append(fld)

def add_hyperlink(paragraph, url, text):
    part=paragraph.part; r_id=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink=OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'),r_id)
    new_run=OxmlElement('w:r'); rPr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),BLUE); rPr.append(color); u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rPr.append(u); new_run.append(rPr); t=OxmlElement('w:t'); t.text=text; new_run.append(t); hyperlink.append(new_run); paragraph._p.append(hyperlink)

def add_rich_text(paragraph, text, size=10.8, color=DARK):
    # handles **bold**, `code`, and simple URLs
    pattern=re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|https?://\S+)')
    pos=0
    for m in pattern.finditer(text):
        if m.start()>pos:
            r=paragraph.add_run(text[pos:m.start()]); set_run_font(r,size=size,color=color)
        tok=m.group(0)
        if tok.startswith('**'):
            r=paragraph.add_run(tok[2:-2]); set_run_font(r,size=size,color=color,bold=True)
        elif tok.startswith('`'):
            r=paragraph.add_run(tok[1:-1]); set_run_font(r,name='Liberation Mono',size=size-0.3,color=NAVY)
        else:
            add_hyperlink(paragraph,tok.rstrip('.,)'),tok.rstrip('.,)'))
            suffix=tok[len(tok.rstrip('.,)')):]
            if suffix:
                r=paragraph.add_run(suffix); set_run_font(r,size=size,color=color)
        pos=m.end()
    if pos<len(text):
        r=paragraph.add_run(text[pos:]); set_run_font(r,size=size,color=color)

def add_para(doc,text,style=None):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    add_rich_text(p,text)
    return p

def add_code(doc, lines):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(6.45)
    c=t.cell(0,0); shade(c,'F4F6F8'); set_cell_margins(c,130,150,130,150)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run('\n'.join(lines)); set_run_font(r,name='Liberation Mono',size=8.5,color=NAVY)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_markdown_table(doc, rows):
    if not rows: return
    cols=len(rows[0]); table=doc.add_table(rows=1, cols=cols); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    widths=[6.45/cols]*cols
    if cols==2: widths=[2.0,4.45]
    elif cols==3: widths=[1.55,2.5,2.4]
    elif cols==4: widths=[1.2,2.4,1.5,1.35]
    hdr=table.rows[0]; set_repeat_table_header(hdr)
    for i,val in enumerate(rows[0]):
        hdr.cells[i].width=Inches(widths[i]); shade(hdr.cells[i],NAVY); set_cell_margins(hdr.cells[i],110,120,110,120); hdr.cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=hdr.cells[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; r=p.add_run(val.strip()); set_run_font(r,size=9,bold=True,color=WHITE)
    for ri,row in enumerate(rows[1:]):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].width=Inches(widths[i]); set_cell_margins(cells[i],100,120,100,120); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==0: shade(cells[i],'F6F9FB')
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); add_rich_text(p,val.strip(),size=8.7)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def add_picture(doc, rel_path, caption):
    path=ROOT/rel_path
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    with Image.open(path) as im:
        w,h=im.size
    maxw=6.65; maxh=7.5
    width=min(maxw, maxh*w/h)
    r=p.add_run(); r.add_picture(str(path),width=Inches(width))
    cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(8)
    rr=cap.add_run(caption); set_run_font(rr,size=8.5,italic=True,color=GRAY)

def parse_md(doc):
    lines=MD.read_text(encoding='utf-8').splitlines()
    # skip first title/subtitle/meta until first H1 section
    i=0; started=False; code=False; code_lines=[]
    while i<len(lines):
        line=lines[i]
        if line.startswith('# 1. Executive summary'): started=True
        if not started: i+=1; continue
        if line.startswith('```'):
            if not code: code=True; code_lines=[]
            else: add_code(doc,code_lines); code=False
            i+=1; continue
        if code:
            code_lines.append(line); i+=1; continue
        if not line.strip() or line.strip()=='---': i+=1; continue
        m_img=re.match(r'!\[(.*?)\]\((.*?)\)',line.strip())
        if m_img:
            add_picture(doc,m_img.group(2),m_img.group(1)); i+=1; continue
        if line.startswith('# '):
            txt=line[2:].strip(); p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.page_break_before=True; p.paragraph_format.keep_with_next=True; r=p.add_run(txt); set_run_font(r,size=19,bold=True,color=NAVY); i+=1; continue
        if line.startswith('## '):
            txt=line[3:].strip(); p=doc.add_paragraph(style='Heading 2'); p.paragraph_format.keep_with_next=True; r=p.add_run(txt); set_run_font(r,size=14,bold=True,color=TEAL); i+=1; continue
        if line.startswith('### '):
            txt=line[4:].strip(); p=doc.add_paragraph(style='Heading 3'); p.paragraph_format.keep_with_next=True; r=p.add_run(txt); set_run_font(r,size=11.5,bold=True,color=BLUE); i+=1; continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|?\s*[-:]+',lines[i+1].strip().lstrip('|')):
            rows=[]
            rows.append([c.strip() for c in line.strip().strip('|').split('|')]); i+=2
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i+=1
            add_markdown_table(doc,rows); continue
        if re.match(r'^- ',line):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.first_line_indent=Inches(-0.15); add_rich_text(p,line[2:].strip()); i+=1; continue
        if re.match(r'^\d+\. ',line):
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.left_indent=Inches(0.32); p.paragraph_format.first_line_indent=Inches(-0.32); add_rich_text(p,line.strip()); i+=1; continue
        # regular paragraph
        add_para(doc,line.strip()); i+=1


doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(0.72); sec.bottom_margin=Inches(0.65); sec.left_margin=Inches(0.85); sec.right_margin=Inches(0.85)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Liberation Sans'; normal.font.size=Pt(10.8); normal.font.color.rgb=RGBColor.from_string(DARK); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
for sname in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
    st=styles[sname]; st.font.name='Liberation Sans'
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'Liberation Sans'); st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'Liberation Sans')
styles['Heading 1'].font.color.rgb=RGBColor.from_string(NAVY); styles['Heading 1'].font.size=Pt(19); styles['Heading 1'].font.bold=True; styles['Heading 1'].paragraph_format.space_before=Pt(0); styles['Heading 1'].paragraph_format.space_after=Pt(8)
styles['Heading 2'].font.color.rgb=RGBColor.from_string(TEAL); styles['Heading 2'].font.size=Pt(14); styles['Heading 2'].font.bold=True; styles['Heading 2'].paragraph_format.space_before=Pt(10); styles['Heading 2'].paragraph_format.space_after=Pt(5)
styles['Heading 3'].font.color.rgb=RGBColor.from_string(BLUE); styles['Heading 3'].font.size=Pt(11.5); styles['Heading 3'].font.bold=True; styles['Heading 3'].paragraph_format.space_before=Pt(7); styles['Heading 3'].paragraph_format.space_after=Pt(4)

# cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(36); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
r=p.add_run('QUANTUM PARKS'); set_run_font(r,size=13,bold=True,color=TEAL)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(22); p.paragraph_format.space_after=Pt(6)
r=p.add_run('ElevenLabs AI Receptionist Platform'); set_run_font(r,size=30,bold=True,color=NAVY)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(22)
r=p.add_run('Production Architecture Pack'); set_run_font(r,size=22,bold=True,color=BLUE)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(26)
r=p.add_run('Control plane • Company knowledge • Agent testing • Canonical call intelligence • Guardrails'); set_run_font(r,size=12,color=GRAY)

# cover callout table
ct=doc.add_table(rows=1,cols=1); ct.alignment=WD_TABLE_ALIGNMENT.CENTER; ct.autofit=False; ct.columns[0].width=Inches(6.5)
c=ct.cell(0,0); shade(c,LIGHT2); set_cell_margins(c,220,220,220,220)
p=c.paragraphs[0]; r=p.add_run('Architecture position'); set_run_font(r,size=10,bold=True,color=TEAL)
p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); add_rich_text(p,'ElevenLabs is the managed live voice receptionist. The Quantum Parks application is the exclusive normal-user control plane and the authoritative knowledge, testing, call-history, governance and intelligence platform.',size=12,color=NAVY)

doc.add_paragraph().paragraph_format.space_after=Pt(18)
meta=doc.add_table(rows=5,cols=2); meta.alignment=WD_TABLE_ALIGNMENT.LEFT; meta.autofit=False
labels=['Version','Status','Prepared for','Architecture authority','Date']
values=['2.0','Implementation baseline','Quantum Parks stakeholders and implementation teams','PRD v2.0 + One-Shot Production Build Prompt v2.0','22 July 2026']
for idx,(lab,val) in enumerate(zip(labels,values)):
    meta.rows[idx].cells[0].width=Inches(1.7); meta.rows[idx].cells[1].width=Inches(4.8)
    shade(meta.rows[idx].cells[0],NAVY); set_cell_margins(meta.rows[idx].cells[0],90,110,90,110); set_cell_margins(meta.rows[idx].cells[1],90,110,90,110)
    r=meta.rows[idx].cells[0].paragraphs[0].add_run(lab); set_run_font(r,size=9,bold=True,color=WHITE)
    r=meta.rows[idx].cells[1].paragraphs[0].add_run(val); set_run_font(r,size=9,color=DARK)

doc.add_page_break()
# document map
p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.page_break_before=False; r=p.add_run('Document map'); set_run_font(r,size=19,bold=True,color=NAVY)
intro=doc.add_paragraph(); add_rich_text(intro,'This pack translates the corrected PRD and one-shot build prompt into an implementation architecture and a machine-enforceable guardrail set. The narrative, diagrams, matrices, ADRs, schemas and repository rules are designed to be used together.')
sections=[
('Executive and boundaries','Sections 1–3','Product split, scope, principles and conformance.'),
('Solution architecture','Sections 4–13','Context, containers, provider integration, agent, knowledge, voice, tests, live tools, post-call intelligence and data.'),
('Trust and operations','Sections 14–24','Analytics, security, hallucination prevention, coding-agent guardrails, APIs, reliability, deployment, observability, CI/CD, QA and runbooks.'),
('Decisions and implementation assets','Section 25 + appendices','ADRs, risks, open decisions, traceability, guardrails and contract schemas.'),
]
t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
widths=[2.0,1.25,3.2]
for i,h in enumerate(['Area','Location','Purpose']):
    t.rows[0].cells[i].width=Inches(widths[i]); shade(t.rows[0].cells[i],NAVY); set_cell_margins(t.rows[0].cells[i]); r=t.rows[0].cells[i].paragraphs[0].add_run(h); set_run_font(r,size=9,bold=True,color=WHITE)
for ri,row in enumerate(sections):
    cells=t.add_row().cells
    for i,val in enumerate(row):
        cells[i].width=Inches(widths[i]); set_cell_margins(cells[i]);
        if ri%2==0: shade(cells[i],'F6F9FB')
        r=cells[i].paragraphs[0].add_run(val); set_run_font(r,size=9,color=DARK,bold=(i==0))

doc.add_page_break()
parse_md(doc)

# footer/header on all sections
for s in doc.sections:
    hp=s.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=hp.add_run('Quantum Parks • ElevenLabs Architecture Pack v2.0'); set_run_font(r,size=8,color=GRAY)
    add_bottom_border(hp,color='D9E2E8',size='4')
    fp=s.footer.paragraphs[0]; add_page_number(fp)

# document properties
props=doc.core_properties; props.title='Quantum Parks ElevenLabs AI Receptionist Platform - Production Architecture Pack v2.0'; props.subject='Production solution architecture and guardrail specification'; props.author='Quantum Parks Architecture'; props.keywords='ElevenLabs, AI receptionist, architecture, knowledge, testing, call intelligence, guardrails'

doc.save(OUT)
print(OUT)
