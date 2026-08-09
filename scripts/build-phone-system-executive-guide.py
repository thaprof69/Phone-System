from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/operations/phone-system-executive-feature-guide.md"
OUT = ROOT / "docs/operations/Quantum_Parks_Phone_System_Executive_Feature_Guide.docx"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(rows[0]):
        set_cell_text(hdr[i], text, bold=True)
        shade_cell(hdr[i], "E8F1FA")
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
    doc.add_paragraph()


def apply_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Heading 1", 18, RGBColor(8, 43, 87)),
        ("Heading 2", 14, RGBColor(24, 83, 70)),
        ("Heading 3", 11.5, RGBColor(60, 60, 60)),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(5)

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(8, 43, 87)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Quantum Parks Phone System")
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(8, 43, 87)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Executive Feature Guide and User Story")
    run.font.name = "Arial"
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(24, 83, 70)

    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout.cell(0, 0)
    shade_cell(cell, "F3F7FB")
    set_cell_text(
        cell,
        "A governed AI receptionist platform that gives Quantum Parks visibility into every call, every referenced policy, every knowledge gap, every follow-up, and every cost signal.",
        bold=True,
    )
    doc.add_section(WD_SECTION_START.NEW_PAGE)


def main():
    text = SOURCE.read_text()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    apply_styles(doc)
    add_cover(doc)

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("# "):
            # Cover already supplies the title.
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("**") and line.endswith("**") and len(line) < 90:
            doc.add_heading(line.strip("*"), level=3)
        elif line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].startswith("| "):
                raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= {"-", " "} for c in raw):
                    rows.append(raw)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:])
        else:
            p = doc.add_paragraph()
            parts = line.split("**")
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                run.bold = idx % 2 == 1
                run.font.name = "Arial"
                run.font.size = Pt(10.5)
        i += 1

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Quantum Parks Phone System Executive Feature Guide")
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 120, 120)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
